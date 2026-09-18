from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

STATUS_FILL = {
    "PASS": "E2F0D9",
    "PASS_NULL": "E2F0D9",
    "PASS_OUTSIDE_HANOI": "E2F0D9",
    "PASS_AMBIGUOUS_0": "E2F0D9",
    "KNOWN_GT_PASS": "E2F0D9",
    "WRONG_LOCATION": "F4CCCC",
    "WRONG_PROVINCE": "F4CCCC",
    "WRONG_STREET": "F4CCCC",
    "REGRESSION": "F4CCCC",
    "KNOWN_BAD_REPEATED": "F4CCCC",
    "MISSING_STREET": "FCE5CD",
    "POTENTIAL_REGRESSION": "FCE5CD",
    "NEED_REVIEW": "FFF2CC",
}


def export_xlsx(df: pd.DataFrame) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "reviewed_data"
    headers = [str(c) for c in df.columns]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in df.itertuples(index=False, name=None):
        ws.append([None if pd.isna(v) else v for v in row])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    status_col = headers.index("final_status") + 1 if "final_status" in headers else None
    if status_col:
        for r in range(2, ws.max_row + 1):
            status = str(ws.cell(r, status_col).value or "")
            if status in STATUS_FILL:
                ws.cell(r, status_col).fill = PatternFill("solid", fgColor=STATUS_FILL[status])
                ws.cell(r, status_col).font = Font(bold=True)

    for idx, header in enumerate(headers, start=1):
        if header in {"final_note", "rule_note", "ai_reason", "geocoder_forward", "geocoder_reverse", "description"}:
            width = 55
        elif header in {"extract_location", "location", "ai_suggested_location", "url"}:
            width = 40
        else:
            width = min(28, max(12, len(header) + 2))
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    summary = wb.create_sheet("summary")
    summary.append(["Status", "Count"])
    counts = df["final_status"].fillna("").astype(str).value_counts()
    for status, count in counts.items():
        summary.append([status, int(count)])
    summary["D1"] = "Rule chính"
    rules = [
        "NULL + 0,0 => PASS_NULL",
        "Ngoài Hà Nội rõ ràng + 0,0 => PASS_OUTSIDE_HANOI",
        "Tên trùng/mơ hồ + không đủ context: AI có thể PASS_AMBIGUOUS_0",
        "Có tên đường nhưng location mất đường => MISSING_STREET",
        "Context Hà Nội rõ + 0,0 => cần kiểm tra regression",
        "Không coi 'trung tâm' mặc định là 'Đường Trung Tâm'",
        "DVHC cũ vẫn có thể đúng nếu cùng địa điểm thực tế",
    ]
    for i, text in enumerate(rules, start=2):
        summary.cell(i, 4).value = text
    for c in summary[1]:
        c.font = Font(bold=True)
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 12
    summary.column_dimensions["D"].width = 80

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
